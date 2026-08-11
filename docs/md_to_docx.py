"""Convert Markdown docs to DOCX (headings, tables, lists, code, paragraphs)."""
from __future__ import annotations

import re
import sys
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor


def set_run_font(run, name: str = "Times New Roman", size: int | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, node):
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            run = paragraph.add_run(text)
            set_run_font(run, size=11)
        return

    if not isinstance(node, Tag):
        return

    name = node.name.lower()
    if name in ("strong", "b"):
        run = paragraph.add_run(node.get_text())
        set_run_font(run, size=11, bold=True)
    elif name in ("em", "i"):
        run = paragraph.add_run(node.get_text())
        set_run_font(run, size=11)
        run.italic = True
    elif name == "code":
        run = paragraph.add_run(node.get_text())
        set_run_font(run, name="Consolas", size=9)
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    elif name == "a":
        run = paragraph.add_run(node.get_text())
        set_run_font(run, size=11)
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        run.underline = True
    elif name == "br":
        paragraph.add_run("\n")
    else:
        for child in node.children:
            add_inline(paragraph, child)


def add_paragraph_from_tag(doc: Document, tag: Tag, style: str | None = None):
    p = doc.add_paragraph(style=style)
    for child in tag.children:
        add_inline(p, child)
    for run in p.runs:
        if run.font.size is None:
            set_run_font(run, size=11)
    return p


def add_table(doc: Document, table_tag: Tag):
    rows = table_tag.find_all("tr")
    if not rows:
        return
    grid = []
    for tr in rows:
        cells = tr.find_all(["th", "td"])
        grid.append([c.get_text(" ", strip=True) for c in cells])
    cols = max(len(r) for r in grid)
    table = doc.add_table(rows=len(grid), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(grid):
        for j in range(cols):
            text = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            is_header = i == 0
            set_run_font(run, size=9, bold=is_header)
    doc.add_paragraph()


def add_code_block(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text.rstrip("\n"))
    set_run_font(run, name="Consolas", size=8)
    # light background via shading
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F4F6")
    shd.set(qn("w:val"), "clear")
    p._p.get_or_add_pPr().append(shd)


def convert_md_to_docx(md_path: Path, docx_path: Path):
    md_text = md_path.read_text(encoding="utf-8")
    html = markdown.markdown(
        md_text,
        extensions=["tables", "fenced_code", "nl2br", "sane_lists"],
    )
    soup = BeautifulSoup(html, "html.parser")
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    for el in soup.children:
        if isinstance(el, NavigableString):
            continue
        if not isinstance(el, Tag):
            continue
        name = el.name.lower()
        if name == "h1":
            p = add_paragraph_from_tag(doc, el)
            for run in p.runs:
                set_run_font(run, size=18, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif name == "h2":
            p = add_paragraph_from_tag(doc, el)
            for run in p.runs:
                set_run_font(run, size=14, bold=True)
                run.font.color.rgb = RGBColor(0x00, 0x20, 0x45)
        elif name == "h3":
            p = add_paragraph_from_tag(doc, el)
            for run in p.runs:
                set_run_font(run, size=12, bold=True)
                run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        elif name == "h4":
            p = add_paragraph_from_tag(doc, el)
            for run in p.runs:
                set_run_font(run, size=11, bold=True)
        elif name == "p":
            add_paragraph_from_tag(doc, el)
        elif name == "ul":
            for li in el.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Bullet")
                for child in li.children:
                    if isinstance(child, Tag) and child.name in ("ul", "ol"):
                        continue
                    add_inline(p, child)
        elif name == "ol":
            for li in el.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Number")
                for child in li.children:
                    if isinstance(child, Tag) and child.name in ("ul", "ol"):
                        continue
                    add_inline(p, child)
        elif name == "table":
            add_table(doc, el)
        elif name == "pre":
            code = el.get_text()
            add_code_block(doc, code)
        elif name == "blockquote":
            p = add_paragraph_from_tag(doc, el)
            p.paragraph_format.left_indent = Inches(0.3)
            for run in p.runs:
                run.italic = True
        elif name == "hr":
            doc.add_paragraph("─" * 40)

    doc.save(docx_path)
    print(f"OK: {docx_path}")


def main():
    docs = Path(__file__).resolve().parent
    files = [
        docs / "CHAPITRE_5_ANALYSE_CONCEPTUELLE.md",
        docs / "CHRONOGRAMME_STAGE.md",
    ]
    for md in files:
        if not md.exists():
            print(f"Missing: {md}", file=sys.stderr)
            continue
        convert_md_to_docx(md, md.with_suffix(".docx"))


if __name__ == "__main__":
    main()
